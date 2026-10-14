"""Output writers — produce .txt, .docx, and .pdf files from redacted text."""

import logging
from pathlib import Path

from .config import claim_name

logger = logging.getLogger("pii_buddy")


def write_txt(text: str, path: Path) -> None:
    """Write plain text output."""
    path.write_text(text, encoding="utf-8")


def write_docx(text: str, path: Path) -> None:
    """Write a simple DOCX file. One paragraph per text block."""
    from docx import Document

    doc = Document()
    for paragraph in text.split("\n"):
        doc.add_paragraph(paragraph)
    doc.save(str(path))


def write_pdf(text: str, path: Path) -> None:
    """Write a simple text-based PDF via fpdf2."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for line in text.split("\n"):
        pdf.cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def write_output(
    text: str,
    path: Path,
    output_format: str,
    input_suffix: str,
    overwrite_existing: bool = False,
) -> Path:
    """
    Dispatcher — write output in the appropriate format.

    Args:
        text: Redacted text content.
        path: Output path (without extension — we'll set the correct one).
        output_format: "txt" or "same".
        input_suffix: Original file's suffix (e.g. ".pdf", ".docx", ".txt").
        overwrite_existing: Replace the input file in place. Honoured only
            when the output extension equals the input's; otherwise (and
            when False) an existing file is never touched and a unique name
            is chosen instead.
    """
    name = path.name
    parent = path.parent

    if output_format == "same" and input_suffix.lower() in (".docx", ".doc"):
        ext, writer = ".docx", write_docx
    elif output_format == "same" and input_suffix.lower() == ".pdf":
        ext, writer = ".pdf", write_pdf
    else:
        ext, writer = ".txt", write_txt

    # Overwriting is only safe when the output is the input file itself. A
    # .doc -> .docx rewrite could otherwise clobber an unrelated sibling.
    if overwrite_existing and ext == input_suffix.lower():
        out = parent / f"{name}{ext}"
    else:
        out = claim_name(parent, name, ext, 0o666)
    writer(text, out)
    return out
